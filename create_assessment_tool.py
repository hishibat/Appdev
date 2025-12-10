#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
IT Governance Quick Assessment Tool Generator
Excel形式のアセスメントツールを生成するスクリプト
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

def create_assessment_excel():
    """ITガバナンスアセスメントツールのExcelファイルを作成"""

    # ワークブック作成
    wb = Workbook()

    # デフォルトシートを削除
    if 'Sheet' in wb.sheetnames:
        wb.remove(wb['Sheet'])

    # ===== Sheet1: Input_Form (顧客用) =====
    ws_input = wb.create_sheet("Input_Form", 0)
    create_input_form(ws_input)

    # ===== Sheet2: Analysis_Internal (コンサル用) =====
    ws_analysis = wb.create_sheet("Analysis_Internal", 1)
    create_analysis_sheet(ws_analysis)

    # ===== Sheet3: DB_AllResponses (全顧客結果DB) =====
    ws_db = wb.create_sheet("DB_AllResponses", 2)
    create_database_sheet(ws_db)

    # ファイル保存
    filename = "IT_Governance_QuickAssessment.xlsx"
    wb.save(filename)
    print(f"✓ {filename} を作成しました")

    return filename

def create_input_form(ws):
    """Input_Formシートを作成（顧客用インターフェース）"""

    # スタイル定義
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    subheader_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    subheader_font = Font(bold=True, size=11)
    result_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    result_font = Font(bold=True, size=11, color="C65911")

    border_thin = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # 列幅設定
    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 50
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 50
    ws.column_dimensions['E'].width = 15

    # ===== ヘッダーセクション =====
    ws['A1'] = "IT Governance Quick Assessment Tool"
    ws['A1'].font = Font(bold=True, size=16, color="366092")
    ws.merge_cells('A1:E1')

    ws['A2'] = "本アセスメントは、貴社のITガバナンス成熟度を簡易的に診断するツールです"
    ws.merge_cells('A2:E2')

    # 入力欄
    ws['A4'] = "会社名:"
    ws['A4'].font = Font(bold=True)
    ws['B4'].fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")
    ws.merge_cells('B4:C4')

    ws['A5'] = "担当者名:"
    ws['A5'].font = Font(bold=True)
    ws['B5'].fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")
    ws.merge_cells('B5:C5')

    ws['A6'] = "実施日:"
    ws['A6'].font = Font(bold=True)
    ws['B6'].fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")
    ws['B6'] = "=TODAY()"
    ws['B6'].number_format = "yyyy/mm/dd"
    ws.merge_cells('B6:C6')

    # ===== 質問セクション =====
    row = 8
    ws[f'A{row}'] = "アセスメント質問"
    ws[f'A{row}'].font = header_font
    ws[f'A{row}'].fill = header_fill
    ws.merge_cells(f'A{row}:C{row}')

    # 質問テーブルヘッダー
    row = 9
    headers = ['No', '質問内容', '回答']
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = subheader_font
        cell.fill = subheader_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border_thin

    # 質問リスト（カテゴリ別）
    questions = [
        # 戦略 (Strategy)
        ("戦略", "ITガバナンスの方針や戦略が明文化され、経営層に承認されている"),
        ("戦略", "IT投資の優先順位付けプロセスが確立され、事業目標と連携している"),
        ("戦略", "ITロードマップが定期的に見直され、最新の事業ニーズを反映している"),
        ("戦略", "IT部門のKPI/目標が事業目標と整合している"),
        ("戦略", "デジタル変革(DX)の戦略と実行計画が策定されている"),

        # 組織 (Organization)
        ("組織", "IT部門の役割・責任が明確に定義されている"),
        ("組織", "IT部門とビジネス部門の協働体制が構築されている"),
        ("組織", "ITガバナンス委員会などの意思決定機関が機能している"),
        ("組織", "IT人材の育成・評価制度が整備されている"),
        ("組織", "重要なIT人材の後継者計画が策定されている"),

        # プロセス (Process)
        ("プロセス", "プロジェクト管理プロセス(PMO機能)が確立している"),
        ("プロセス", "IT予算管理プロセスが適切に運用されている"),
        ("プロセス", "ベンダー管理プロセスが確立している"),
        ("プロセス", "IT資産管理プロセスが整備されている"),
        ("プロセス", "変更管理プロセスが適切に機能している"),

        # リスク・コンプライアンス (Risk & Compliance)
        ("リスク", "ITリスク評価が定期的に実施されている"),
        ("リスク", "情報セキュリティポリシーが策定され、周知されている"),
        ("リスク", "事業継続計画(BCP/DR)が策定され、定期的にテストされている"),
        ("リスク", "コンプライアンス要件(個人情報保護法等)への対応が適切に実施されている"),
        ("リスク", "インシデント管理プロセスが確立している"),

        # パフォーマンス管理 (Performance)
        ("パフォーマンス", "ITサービスレベル(SLA)が定義され、モニタリングされている"),
        ("パフォーマンス", "システムパフォーマンスが定期的に測定・報告されている"),
        ("パフォーマンス", "IT投資効果(ROI)の測定と評価が実施されている"),
        ("パフォーマンス", "利用者満足度調査が定期的に実施されている"),
        ("パフォーマンス", "ITサービス改善の取り組みが継続的に行われている"),
    ]

    # プルダウンの選択肢
    dv = DataValidation(type="list", formula1='"Level 1,Level 2,Level 3,Level 4"', allow_blank=True)
    dv.prompt = "レベルを選択してください"
    dv.promptTitle = "回答選択"
    ws.add_data_validation(dv)

    # 質問を配置
    for idx, (category, question) in enumerate(questions, start=1):
        row = 9 + idx
        ws[f'A{row}'] = idx
        ws[f'B{row}'] = question
        ws[f'C{row}'] = ""  # 回答欄

        # スタイル適用
        ws[f'A{row}'].alignment = Alignment(horizontal='center', vertical='center')
        ws[f'B{row}'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        ws[f'C{row}'].alignment = Alignment(horizontal='center', vertical='center')

        for col in ['A', 'B', 'C']:
            ws[f'{col}{row}'].border = border_thin

        # プルダウン追加
        dv.add(f'C{row}')

        # 行の高さを調整
        ws.row_dimensions[row].height = 30

    last_question_row = 9 + len(questions)

    # ===== 集計ボタン説明 =====
    row = last_question_row + 2
    ws[f'A{row}'] = "※ 全ての質問に回答後、VBAマクロの「集計ボタン」を実行してください"
    ws[f'A{row}'].font = Font(italic=True, color="C65911")
    ws.merge_cells(f'A{row}:C{row}')

    # ===== 結果表示エリア =====
    row = last_question_row + 4
    ws[f'D{row}'] = "【 簡易診断結果 】"
    ws[f'D{row}'].font = result_font
    ws[f'D{row}'].fill = result_fill
    ws.merge_cells(f'D{row}:E{row}')

    result_start_row = row + 1

    # 総合スコア
    ws[f'D{result_start_row}'] = "総合スコア:"
    ws[f'D{result_start_row}'].font = Font(bold=True)
    ws[f'E{result_start_row}'] = "=Analysis_Internal!B2"
    ws[f'E{result_start_row}'].number_format = "0.00"
    ws[f'E{result_start_row}'].alignment = Alignment(horizontal='center')

    # カテゴリ別スコア
    categories = ["戦略", "組織", "プロセス", "リスク", "パフォーマンス"]
    for idx, cat in enumerate(categories, start=1):
        current_row = result_start_row + idx
        ws[f'D{current_row}'] = f"{cat}:"
        ws[f'E{current_row}'] = f"=Analysis_Internal!B{2 + idx}"
        ws[f'E{current_row}'].number_format = "0.00"
        ws[f'E{current_row}'].alignment = Alignment(horizontal='center')

    # 評価コメント
    comment_row = result_start_row + len(categories) + 2
    ws[f'D{comment_row}'] = "【 総合評価 】"
    ws[f'D{comment_row}'].font = Font(bold=True)
    ws.merge_cells(f'D{comment_row}:E{comment_row}')

    ws[f'D{comment_row + 1}'] = "=Analysis_Internal!B10"
    ws.merge_cells(f'D{comment_row + 1}:E{comment_row + 3}')
    ws[f'D{comment_row + 1}'].alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)

    print("✓ Input_Form シートを作成しました")

def create_analysis_sheet(ws):
    """Analysis_Internalシートを作成（コンサル用分析）"""

    # スタイル定義
    header_fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)

    # 列幅設定
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 50

    # ===== ヘッダー =====
    ws['A1'] = "内部分析シート（コンサル専用）"
    ws['A1'].font = Font(bold=True, size=14, color="70AD47")
    ws.merge_cells('A1:C1')

    # ===== 集計結果エリア =====
    ws['A2'] = "総合スコア"
    ws['A2'].font = header_font
    ws['A2'].fill = header_fill
    ws['B2'] = "=AVERAGE(B3:B7)"
    ws['B2'].number_format = "0.00"
    ws['C2'] = "全カテゴリの平均値"

    # カテゴリ別スコア
    categories = [
        ("戦略スコア", "B3", "戦略領域の質問（Q1-Q5）の平均"),
        ("組織スコア", "B4", "組織領域の質問（Q6-Q10）の平均"),
        ("プロセススコア", "B5", "プロセス領域の質問（Q11-Q15）の平均"),
        ("リスクスコア", "B6", "リスク領域の質問（Q16-Q20）の平均"),
        ("パフォーマンススコア", "B7", "パフォーマンス領域の質問（Q21-Q25）の平均"),
    ]

    for idx, (cat_name, cell, description) in enumerate(categories, start=3):
        ws[f'A{idx}'] = cat_name
        ws[f'A{idx}'].font = Font(bold=True)

        # 各カテゴリの計算式（5問ずつ）
        start_q = (idx - 3) * 5 + 10  # Input_Formの質問開始行
        end_q = start_q + 4

        # Level 1-4を1-4の数値に変換して平均
        formula = f'=AVERAGE('
        for q in range(start_q, end_q + 1):
            formula += f'IF(Input_Form!C{q}="Level 1",1,IF(Input_Form!C{q}="Level 2",2,IF(Input_Form!C{q}="Level 3",3,IF(Input_Form!C{q}="Level 4",4,0)))),'
        formula = formula.rstrip(',') + ')'

        ws[cell] = formula
        ws[cell].number_format = "0.00"
        ws[f'C{idx}'] = description

    # ===== 評価コメント生成 =====
    ws['A9'] = "評価コメント"
    ws['A9'].font = header_font
    ws['A9'].fill = header_fill

    comment_formula = '''=IF(B2>=3.5,"【優良レベル】ITガバナンス体制は十分に整備されています。継続的な改善を推奨します。",IF(B2>=2.5,"【標準レベル】ITガバナンスの基本は整っていますが、一部改善の余地があります。特に低スコアの領域に注目してください。",IF(B2>=1.5,"【改善必要】ITガバナンス体制に大きな改善が必要です。戦略的な取り組みをお勧めします。","【早急対応】ITガバナンスが未整備の状態です。早急な体制構築が必要です。")))'''

    ws['B10'] = comment_formula
    ws.merge_cells('B10:C12')
    ws['B10'].alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)

    # ===== 詳細分析エリア =====
    row = 14
    ws[f'A{row}'] = "詳細分析データ"
    ws[f'A{row}'].font = header_font
    ws[f'A{row}'].fill = header_fill
    ws.merge_cells(f'A{row}:C{row}')

    row += 1
    ws[f'A{row}'] = "質問No"
    ws[f'B{row}'] = "回答値"
    ws[f'C{row}'] = "回答テキスト"

    for col in ['A', 'B', 'C']:
        ws[f'{col}{row}'].font = Font(bold=True)

    # 各質問の回答を数値化
    for q_num in range(1, 26):  # 25問
        row += 1
        input_row = 9 + q_num
        ws[f'A{row}'] = f"Q{q_num}"
        ws[f'B{row}'] = f'=IF(Input_Form!C{input_row}="Level 1",1,IF(Input_Form!C{input_row}="Level 2",2,IF(Input_Form!C{input_row}="Level 3",3,IF(Input_Form!C{input_row}="Level 4",4,0))))'
        ws[f'C{row}'] = f"=Input_Form!C{input_row}"

    print("✓ Analysis_Internal シートを作成しました")

def create_database_sheet(ws):
    """DB_AllResponsesシートを作成（全顧客結果DB）"""

    # スタイル定義
    header_fill = PatternFill(start_color="C65911", end_color="C65911", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)

    # 列幅設定
    columns = ['日付', '会社名', '担当者', '総合スコア', '戦略', '組織', 'プロセス', 'リスク', 'パフォーマンス']
    for idx, col_name in enumerate(columns, start=1):
        col_letter = get_column_letter(idx)
        ws.column_dimensions[col_letter].width = 15

    # ヘッダー行
    for idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # サンプルデータ（オプション）
    ws['A2'] = "※ VBAマクロで自動追記されます"
    ws.merge_cells('A2:I2')

    print("✓ DB_AllResponses シートを作成しました")

def create_vba_code():
    """VBAマクロコードを生成"""

    vba_code = '''Sub CalculateAssessment()
    '
    ' IT Governance Assessment 集計マクロ
    '
    Dim wsInput As Worksheet
    Dim wsAnalysis As Worksheet
    Dim wsDB As Worksheet
    Dim lastRow As Long
    Dim i As Integer
    Dim allAnswered As Boolean
    Dim companyName As String
    Dim personName As String

    ' シート参照
    Set wsInput = ThisWorkbook.Sheets("Input_Form")
    Set wsAnalysis = ThisWorkbook.Sheets("Analysis_Internal")
    Set wsDB = ThisWorkbook.Sheets("DB_AllResponses")

    ' 入力チェック
    allAnswered = True
    For i = 10 To 34  ' 質問行（25問）
        If wsInput.Range("C" & i).Value = "" Then
            allAnswered = False
            Exit For
        End If
    Next i

    If Not allAnswered Then
        MsgBox "全ての質問に回答してください。", vbExclamation, "入力エラー"
        Exit Sub
    End If

    ' 会社名・担当者名チェック
    companyName = wsInput.Range("B4").Value
    personName = wsInput.Range("B5").Value

    If companyName = "" Or personName = "" Then
        MsgBox "会社名と担当者名を入力してください。", vbExclamation, "入力エラー"
        Exit Sub
    End If

    ' 強制再計算
    Application.Calculate

    ' 結果をDBシートに追記（オプション）
    lastRow = wsDB.Cells(wsDB.Rows.Count, 1).End(xlUp).Row + 1

    wsDB.Cells(lastRow, 1).Value = Date  ' 日付
    wsDB.Cells(lastRow, 2).Value = companyName  ' 会社名
    wsDB.Cells(lastRow, 3).Value = personName  ' 担当者名
    wsDB.Cells(lastRow, 4).Value = wsAnalysis.Range("B2").Value  ' 総合スコア
    wsDB.Cells(lastRow, 5).Value = wsAnalysis.Range("B3").Value  ' 戦略
    wsDB.Cells(lastRow, 6).Value = wsAnalysis.Range("B4").Value  ' 組織
    wsDB.Cells(lastRow, 7).Value = wsAnalysis.Range("B5").Value  ' プロセス
    wsDB.Cells(lastRow, 8).Value = wsAnalysis.Range("B6").Value  ' リスク
    wsDB.Cells(lastRow, 9).Value = wsAnalysis.Range("B7").Value  ' パフォーマンス

    ' 完了メッセージ
    MsgBox "集計が完了しました！" & vbCrLf & vbCrLf & _
           "総合スコア: " & Format(wsAnalysis.Range("B2").Value, "0.00") & vbCrLf & _
           "結果は右側の【簡易診断結果】エリアに表示されています。", _
           vbInformation, "集計完了"

End Sub

Sub ProtectAnalysisSheet()
    '
    ' Analysis_Internal シートを保護（VeryHidden設定）
    '
    Dim ws As Worksheet
    Set ws = ThisWorkbook.Sheets("Analysis_Internal")

    ' シートを非常に非表示に設定
    ws.Visible = xlSheetVeryHidden

    MsgBox "Analysis_Internal シートを非表示にしました。", vbInformation
End Sub

Sub UnprotectAnalysisSheet()
    '
    ' Analysis_Internal シートの保護を解除
    ' ※パスワード: "abeam2025"
    '
    Dim password As String
    Dim inputPwd As String
    Dim ws As Worksheet

    password = "abeam2025"
    inputPwd = InputBox("パスワードを入力してください:", "シート表示")

    If inputPwd = password Then
        Set ws = ThisWorkbook.Sheets("Analysis_Internal")
        ws.Visible = xlSheetVisible
        MsgBox "Analysis_Internal シートを表示しました。", vbInformation
    Else
        MsgBox "パスワードが正しくありません。", vbCritical
    End If
End Sub
'''

    # VBAコードをファイルに保存
    with open("VBA_Macro_Code.bas", "w", encoding="utf-8") as f:
        f.write(vba_code)

    print("✓ VBAマクロコードを VBA_Macro_Code.bas に保存しました")

def create_instruction_file():
    """使用方法の説明ファイルを作成"""

    instructions = """# IT Governance Quick Assessment Tool - 使用方法

## 📋 概要
このExcelツールは、ITガバナンス成熟度を簡易診断するためのアセスメントツールです。

## 📁 ファイル構成
- `IT_Governance_QuickAssessment.xlsx` - メインのExcelファイル
- `VBA_Macro_Code.bas` - VBAマクロコード
- `使用方法.md` - このファイル

## 🔧 初期設定（マクロの追加）

### 1. Excelファイルを.xlsmとして保存
1. `IT_Governance_QuickAssessment.xlsx` を開く
2. 「名前を付けて保存」→ 「Excelマクロ有効ブック (.xlsm)」を選択
3. ファイル名を `IT_Governance_QuickAssessment.xlsm` に変更して保存

### 2. VBAマクロをインポート
1. Excel上で `Alt + F11` を押してVBAエディタを開く
2. メニューから「ファイル」→「ファイルのインポート」を選択
3. `VBA_Macro_Code.bas` を選択してインポート
4. VBAエディタを閉じる

### 3. マクロボタンの配置
1. Input_Formシートを開く
2. 「開発」タブ →「挿入」→「ボタン（フォームコントロール）」を選択
3. 質問リストの下にボタンを配置
4. マクロ選択画面で `CalculateAssessment` を選択
5. ボタンのテキストを「集計する」に変更

## 📝 使用方法

### 顧客側の操作
1. **Input_Formシート**を開く
2. **会社名**、**担当者名**を入力
3. 25の質問に対して、プルダウンから**Level 1〜4**を選択
   - Level 1: 未整備・未実施
   - Level 2: 一部実施・改善必要
   - Level 3: 概ね実施・標準レベル
   - Level 4: 十分実施・優良レベル
4. 全て回答後、**「集計する」ボタン**をクリック
5. 右側の「簡易診断結果」エリアに結果が表示される
6. ファイルを保存して返送

### コンサルタント側の操作
1. 顧客から返送されたファイルを開く
2. `Alt + F8` でマクロ一覧を開き、`UnprotectAnalysisSheet` を実行
3. パスワード: `abeam2025` を入力
4. **Analysis_Internal**シートが表示される
5. 詳細なスコア・分析データを確認
6. 必要に応じて**DB_AllResponses**シートで複数顧客の結果を比較

## 📊 シート説明

### Sheet1: Input_Form（顧客用）
- 質問への回答入力
- 簡易診断結果の表示
- 顧客に見せる唯一のシート

### Sheet2: Analysis_Internal（コンサル専用）
- 詳細なスコア計算
- カテゴリ別分析
- 評価コメント自動生成
- 通常は非表示

### Sheet3: DB_AllResponses（データベース）
- 全顧客の結果を蓄積
- マクロ実行時に自動追記
- 横断分析用

## 🔒 セキュリティ設定

### Analysis_Internalシートを非表示にする
1. `Alt + F8` でマクロ実行
2. `ProtectAnalysisSheet` を実行
3. シートが非常に非表示になる（通常の方法では表示不可）

### 顧客配布時のチェックリスト
- [ ] Analysis_Internalシートが非常に非表示になっている
- [ ] DB_AllResponsesシートが非表示になっている
- [ ] Input_Formシートのみ表示されている
- [ ] マクロが有効になっている
- [ ] 集計ボタンが正しく動作する

## 💡 カスタマイズ

### 質問の追加・変更
1. Input_Formシートの質問リストを編集
2. Analysis_Internalシートの集計式を調整
3. カテゴリ数や配分を変更する場合は数式も調整

### スコア計算ロジックの変更
- Analysis_Internalシートのセル `B3:B7` の数式を編集
- 重み付けを追加する場合は SUMPRODUCT 関数を使用

### 評価コメントのカスタマイズ
- Analysis_Internalシートのセル `B10` の IF 文を編集
- スコア閾値や評価文言を変更可能

## 🆘 トラブルシューティング

### マクロが実行できない
- Excelのセキュリティ設定を確認
- 「ファイル」→「オプション」→「セキュリティセンター」→「マクロの設定」
- 「警告を表示してすべてのマクロを無効にする」を選択

### 結果が表示されない
- 数式エラーがないか確認
- `F9` キーで強制再計算を実行
- Analysis_Internalシートの数式を確認

### パスワードを忘れた
- VBAエディタで `UnprotectAnalysisSheet` マクロを編集
- `password = "abeam2025"` の行でパスワードを確認

## 📞 サポート
ご不明点があれば、開発者にお問い合わせください。

---
© 2025 IT Governance Assessment Tool
"""

    with open("使用方法.md", "w", encoding="utf-8") as f:
        f.write(instructions)

    print("✓ 使用方法.md を作成しました")

if __name__ == "__main__":
    print("=" * 60)
    print("IT Governance Quick Assessment Tool Generator")
    print("=" * 60)
    print()

    # Excelファイル作成
    create_assessment_excel()

    # VBAマクロコード生成
    create_vba_code()

    # 使用方法ドキュメント生成
    create_instruction_file()

    print()
    print("=" * 60)
    print("✅ 全てのファイルが正常に生成されました！")
    print("=" * 60)
    print()
    print("生成されたファイル:")
    print("  1. IT_Governance_QuickAssessment.xlsx")
    print("  2. VBA_Macro_Code.bas")
    print("  3. 使用方法.md")
    print()
    print("次のステップ:")
    print("  1. .xlsxファイルを.xlsmとして保存")
    print("  2. VBAマクロをインポート")
    print("  3. 集計ボタンを配置")
    print()
    print("詳細は「使用方法.md」をご覧ください。")
    print("=" * 60)
